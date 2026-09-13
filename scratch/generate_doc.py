import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

doc = Document()

# Set Margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Base Style Settings
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = RGBColor(0x2B, 0x2D, 0x42) # Slate Dark

# Helper functions for adding elements
def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy
    p.paragraph_format.space_after = Pt(4)
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63) # Neutral Grey
    p.paragraph_format.space_after = Pt(24)
    return p

def add_heading_1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_bullet(p_text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    r_text = p.add_run(p_text)
    r_text.font.color.rgb = RGBColor(0x2B, 0x2D, 0x42)
    return p

def add_callout(text, title=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F3F4F6") # Light grey background
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if title:
        r_title = p.add_run(title + "\n")
        r_title.font.bold = True
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(10.5)
    r_text.font.italic = True
    r_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# --- DOCUMENT CONTENT GENERATION ---

add_title("PROJECT REQUIREMENTS & TECHNICAL OVERVIEW")
add_subtitle("International Conference Web Platform (ICSWHMH)\nRajagiri College of Social Sciences")

# Section 1: Executive Summary
add_heading_1("1. Executive Summary")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
p.add_run(
    "The International Conference Web Platform is a comprehensive digital solution designed for Rajagiri College of Social Sciences "
    "to manage and host the International Conference on Social Work, Health, and Mental Health (ICSWHMH). The system serves as the primary "
    "digital gateway for international delegates, researchers, keynote speakers, and organizers. It bridges public-facing informational modules "
    "with secure administrative back-end workflows, enabling seamless registration, abstract submissions, venue details, payment processing, and event administration."
)

# Section 2: Current Problem
add_heading_1("2. Current Problem Statement")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
p.add_run(
    "Organizing large-scale international academic conferences poses severe operational and technological challenges when managed manually or via fragmented tools:"
)

add_bullet(" High risk of data mismatches, lost submission files, and manual entry errors during paper registrations.", "Manual Delegate & Abstract Management: ")
add_bullet(" Delay in payment verification for international and domestic currencies, leading to reconciliation conflicts.", "Unintegrated Payment Gateways: ")
add_bullet(" International attendees lack a single centralized hub to view conference themes, schedules, hotel accommodations, and local tourist attraction details.", "Fragmented Information Delivery: ")
add_bullet(" Organizers lack real-time visibility into registration numbers, payment statuses, and delegate categorizations necessary for logistics planning.", "Lack of Administrative Dashboard & Insights: ")

# Section 3: Objectives
add_heading_1("3. Core Objectives")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
p.add_run("The primary objectives of the project are as follows:")

add_bullet(" Provide a responsive, high-performance web platform detailing conference schedules, keynotes, themes, topics, and venue guides.", "Centralized Information Hub: ")
add_bullet(" Offer an intuitive web form allowing attendees to submit paper abstracts, personal details, and delegate classifications.", "Digital Abstract Submission & Registration: ")
add_bullet(" Seamlessly integrate Razorpay API to process real-time registration fee payments securely with automated tracking.", "Automated Payment Integration: ")
add_bullet(" Provide administrators with a secure dashboard to monitor live registration statistics, verify candidate records, update settings, and export attendee data.", "Robust Administrative Backend: ")
add_bullet(" Ensure data privacy, robust environment configuration handling, cross-browser compatibility, and mobile responsiveness.", "Security & Scalability: ")

# Section 4: Scope of Work
add_heading_1("4. Scope of Work")

add_heading_2("4.1 In-Scope Features")
add_bullet("Public Landing Page (index.html), Themes & Sub-themes (themes.html, topics.html), Abstract Guidelines (abstract.html), Committee details, Tourist Attractions, and Hotel Guides.", "1. Informational & Marketing Portal: ")
add_bullet("Dynamic registration forms (registration.html/php) with live email verification (check_email.php) and registration database storage (save_registration.php).", "2. Delegate Registration System: ")
add_bullet("Integration with Razorpay SDK (process_payment.php, razorpay/ folder) supporting multi-currency/tier payments and transaction verification (view_transaction.php).", "3. Payment Processing Engine: ")
add_bullet("Protected management area (admin/) for viewing registration lists, filtering delegates, managing configuration variables, and exporting data.", "4. Admin Control Panel: ")
add_bullet("Centralized configuration management using environment variables (.env / .env.example) and structured database architecture (backend/db.php).", "5. Core Infrastructure: ")

add_heading_2("4.2 Out-of-Scope (Future Enhancements)")
add_bullet("Automated peer-review assignment engine for scientific committee members.", "1. Full Peer-Review Workflow: ")
add_bullet("In-app live streaming or virtual conference video room integrations.", "2. Virtual Conference Video Streaming: ")
add_bullet("Automated generation and emailing of conference participation certificates.", "3. Automated Certificate Generation: ")

# Section 5: Technology Stack
add_heading_1("5. Technology Stack")

# Tech Stack Table
table = doc.add_table(rows=6, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'

headers = ["Layer / Component", "Technology Used", "Purpose / Rationale"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    set_cell_background(cell, "1E3A8A") # Navy header
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    p = cell.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

data = [
    ("Frontend UI / UX", "HTML5, CSS3, JavaScript (Vanilla ES6), FontAwesome", "Responsive, fast-loading, dynamic user interface without heavy client framework overhead."),
    ("Backend Scripting", "PHP 8.x (Modular Scripts)", "Lightweight, reliable server-side processing for form handling, payment webhooks, and DB interactions."),
    ("Database Management", "MySQL / MariaDB (PDO Driver)", "Relational storage for delegates, payments, and abstracts with prepared statements against SQL injection."),
    ("Payment Gateway", "Razorpay Payment Gateway API & SDK", "Secure processing of domestic and international credit/debit card, UPI, and net-banking transactions."),
    ("Environment & Security", "Dotenv (.env), PHP PDO Prepared Statements", "Secure management of API credentials, database keys, and mitigation of security vulnerabilities.")
]

for row_idx, row_data in enumerate(data, start=1):
    bg_color = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
    for col_idx, text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# Professional Note Box
add_callout(
    "Note on Client & Project Confidentiality:\n"
    "This document summarizes the technical architecture, project scope, and functional objectives for academic and internal evaluation purposes. "
    "Confidential production credentials (such as database passwords, secret keys, or live API keys) are safely isolated in local .env environment configurations and strictly excluded from shared reports.",
    "CLIENT PROJECT GUIDELINES & CONFIDENTIALITY"
)

output_path = "/home/krisb/Desktop/Projects/INTERNATIONAL-CONFERENCE-RAJAGIRI/INTERNATIONAL-CONFERENCE-RAJAGIRI/Executive_Summary_and_Technical_Overview.docx"
doc.save(output_path)
print(f"Document successfully created at: {output_path}")
